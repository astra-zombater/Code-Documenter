import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor

# Native file explorer support
try:
    import tkinter as tk
    from tkinter import filedialog
    HAS_TK = True
except ImportError:
    HAS_TK = False


def add_file_to_doc(doc, file_path, heading_text):
    """Helper to append code content to the document with monospace styling."""
    doc.add_heading(heading_text, level=2)
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            code_content = f.read()

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(12)
        
        run = p.add_run(code_content)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 30, 30)
        
    except Exception as e:
        doc.add_paragraph(f"[Error reading file: {e}]")


def generate_code_docs(targets, doc_filename="Project_Documentation.docx", is_folder_mode=True):
    doc = Document()
    doc.add_heading('Project Source Code Documentation', level=0)
    
    target_extensions = ('.py', '.json', '.yaml', '.yml', '.md', '.sql', '.html', '.css')
    ignore_dirs = {'venv', '.venv', '__pycache__', '.git', 'build', 'dist', '.idea', '.vscode'}

    if is_folder_mode:
        project_dir = Path(targets)
        # Save output document directly inside the selected folder
        output_docx = project_dir / doc_filename
        
        for root, dirs, files in os.walk(project_dir):
            # Exclude virtual environments & internal folders
            dirs[:] = [d for d in dirs if d not in ignore_dirs]
            
            for file in sorted(files):
                if file.endswith(target_extensions):
                    full_path = Path(root) / file
                    rel_path = full_path.relative_to(project_dir)
                    add_file_to_doc(doc, full_path, f"File: {rel_path}")
    else:
        # User selected specific files directly
        files_list = [Path(p) for p in targets]
        # Save output in the parent directory of the first selected file
        output_dir = files_list[0].parent if files_list else Path.cwd()
        output_docx = output_dir / doc_filename
        
        for file_path in sorted(files_list):
            if file_path.is_file():
                add_file_to_doc(doc, file_path, f"File: {file_path.name}")

    doc.save(output_docx)
    print(f"\nDone! Documentation successfully saved to:\n -> {output_docx}")


def pick_folder():
    """Opens a visual folder selection dialog."""
    if HAS_TK:
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        folder = filedialog.askdirectory(title="Select Project Folder")
        root.destroy()
        return folder
    return input("Enter full directory path: ").strip()


def pick_files():
    """Opens a visual file selection dialog."""
    if HAS_TK:
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        files = filedialog.askopenfilenames(
            title="Select Specific Files to Document",
            filetypes=[("Code & Config Files", "*.py *.json *.yaml *.yml *.md *.sql *.html *.css"), ("All Files", "*.*")]
        )
        root.destroy()
        return list(files)
    
    raw = input("Enter full path(s) separated by commas: ").strip()
    return [p.strip() for p in raw.split(",") if p.strip()]


def main():
    print("=" * 50)
    print("       PYTHON CODE DOCUMENTATION GENERATOR")
    print("=" * 50)
    print("1. Document current directory (Default)")
    print("2. Choose a specific folder...")
    print("3. Select specific file(s)...")
    
    choice = input("\nEnter choice (1/2/3, default is 1): ").strip()

    if choice == '2':
        folder_path = pick_folder()
        if folder_path and os.path.exists(folder_path):
            print(f"\nProcessing folder: {folder_path}")
            generate_code_docs(folder_path, is_folder_mode=True)
        else:
            print("No valid folder selected. Exiting.")

    elif choice == '3':
        selected_files = pick_files()
        if selected_files:
            print(f"\nProcessing {len(selected_files)} file(s)...")
            generate_code_docs(selected_files, is_folder_mode=False)
        else:
            print("No files selected. Exiting.")

    else:
        print("\nDocumenting current directory...")
        generate_code_docs(Path.cwd(), is_folder_mode=True)


if __name__ == "__main__":
    main()